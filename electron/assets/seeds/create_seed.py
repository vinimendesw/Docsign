from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document()

title = doc.add_heading('', 0)
title_run = title.add_run('SOLICITACAO DE FERIAS')
title_run.bold = True
title_run.font.size = Pt(16)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('')

body = doc.add_paragraph()
body.add_run('Eu, ')
body.add_run('{{nome_funcionario}}').bold = True
body.add_run(', matricula ')
body.add_run('{{matricula}}').bold = True
body.add_run(', lotado(a) no setor ')
body.add_run('{{setor}}').bold = True
body.add_run(', venho por meio deste solicitar o gozo de ferias regulamentares no periodo de ')
body.add_run('{{data_inicio}}').bold = True
body.add_run(' a ')
body.add_run('{{data_fim}}').bold = True
body.add_run(', totalizando ')
body.add_run('{{dias}}').bold = True
body.add_run(' dias.')

doc.add_paragraph('')

adiantamento = doc.add_paragraph()
adiantamento.add_run('Adiantamento de 13o salario: ')
adiantamento.add_run('{{adiantamento_13}}').bold = True

doc.add_paragraph('')
doc.add_paragraph('Data: ___/___/______')
doc.add_paragraph('')
doc.add_paragraph('Assinatura do Funcionario: _________________________________')
doc.add_paragraph('')
doc.add_paragraph('Assinatura do Gestor: _________________________________')

out = os.path.join(os.path.dirname(__file__), 'ferias_exemplo.docx')
doc.save(out)
print('Criado:', out)
